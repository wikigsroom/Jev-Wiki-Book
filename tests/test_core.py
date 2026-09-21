from __future__ import annotations

import io
import threading
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from PIL import Image

from jev_core.common import query_tokens, tokens
from jev_core.parsing import DocumentParser, passages
from jev_core.retrieval import RagService, Retriever, select_evidence
from jev_core.storage import Cancelled, KnowledgeStore


@pytest.fixture
def store(tmp_path):
    root = tmp_path / "source"
    root.mkdir()
    return KnowledgeStore(tmp_path / "storage", root, DocumentParser())


def build(store, root=None, **kwargs):
    return store.build(root or store.default_root, lambda fields: None, threading.Event(), **kwargs)


def test_docx_reading_order_real_paragraphs_nested_table_and_footer(tmp_path):
    path = tmp_path / "原始文档.docx"
    doc = Document()
    doc.add_paragraph("标题", "Heading 1")
    doc.add_paragraph("")
    doc.add_paragraph("第 3 段保留具体事实")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目甲"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "嵌套数据不能遗漏"
    doc.add_paragraph("表后的第 4 段")
    doc.sections[0].footer.paragraphs[0].text = "页脚内容"
    doc.save(path)
    parsed = DocumentParser().parse(path, path.name)
    blocks = parsed["blocks"]
    assert [b["kind"] for b in blocks] == ["paragraph", "paragraph", "table_row", "paragraph", "paragraph"]
    assert blocks[1]["locator"]["paragraph_number"] == 3
    assert "嵌套数据不能遗漏" in blocks[2]["text"]
    assert "paragraph_number" not in blocks[2]["locator"]
    assert blocks[3]["locator"]["paragraph_number"] == 4
    assert blocks[4]["locator"]["label"] == "页脚第 1 段"


def test_image_occurrences_are_in_order_and_missing_ocr_is_visible(tmp_path):
    path = tmp_path / "有截图.docx"
    doc = Document()
    image = io.BytesIO()
    Image.new("RGB", (12, 12), "white").save(image, "PNG")
    image.seek(0)
    doc.add_paragraph("前文")
    doc.add_picture(image)
    doc.add_paragraph("后文")
    doc.save(path)
    result = DocumentParser().parse(path, path.name)
    assert [b["kind"] for b in result["blocks"]] == ["paragraph", "image", "paragraph"]
    assert result["blocks"][1]["locator"]["paragraph_number"] == 2
    assert result["coverage"]["images_unread"] == 1


def test_word_compatibility_fallback_does_not_duplicate_textboxes(tmp_path):
    from lxml import etree
    from docx.oxml import OxmlElement
    doc = Document()
    para = doc.add_paragraph()
    mc = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    alternate = etree.Element("{" + mc + "}AlternateContent")
    for name in ("Choice", "Fallback"):
        branch = etree.SubElement(alternate, "{" + mc + "}" + name)
        shape, body, p, run, text = [OxmlElement(tag) for tag in ("w:r", "w:txbxContent", "w:p", "w:r", "w:t")]
        text.text = "同一个文本框只索引一次"
        run.append(text); p.append(run); body.append(p); shape.append(body); branch.append(shape)
    para._p.append(alternate)
    path = tmp_path / "fallback.docx"
    doc.save(path)
    blocks = DocumentParser().parse(path, path.name)["blocks"]
    assert [b["text"] for b in blocks] == ["同一个文本框只索引一次"]


def test_long_paragraph_ranges_reconstruct_every_character():
    text = ("完整原文不能截断。" * 200) + "最后的关键数字 24681357"
    block = {"id": "source-p1", "text": text, "locator": {"paragraph_number": 7}}
    chunks = passages(block, 800)
    assert "".join(c["text"] for c in chunks) == text
    assert all(c["text"] == text[c["char_start"]:c["char_end"]] for c in chunks)
    assert all(c["locator"]["paragraph_number"] == 7 for c in chunks)
    assert all(a["char_end"] == b["char_start"] for a, b in zip(chunks, chunks[1:]))
    assert chunks[-1]["text"].endswith("24681357")


def test_pdf_has_real_page_and_bbox_not_fake_paragraph_numbers(tmp_path):
    import fitz
    path = tmp_path / "page.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((50, 60), "first page")
    doc.new_page().insert_text((50, 60), "second page 12345")
    doc.save(path)
    doc.close()
    blocks = DocumentParser().parse(path, path.name)["blocks"]
    assert blocks[1]["locator"]["page"] == 2
    assert len(blocks[1]["locator"]["bbox"]) == 4
    assert "paragraph_number" not in blocks[1]["locator"]


def test_libraries_keep_uploads_and_versions_separate(store, tmp_path):
    (store.default_root / "甲 原文.txt").write_text("每日收益六位小数", encoding="utf-8")
    first = build(store)
    build(store, uploads=[("上传甲.txt", "上传原文".encode())])
    saved = store.snapshot()
    other = tmp_path / "other"
    other.mkdir()
    (other / "乙.txt").write_text("股票行情更新", encoding="utf-8")
    build(store, other)
    assert [d["name"] for d in store.snapshot()["documents"]] == ["乙.txt"]
    build(store, store.default_root)
    assert {d["name"] for d in store.snapshot()["documents"]} == {"甲 原文.txt", "上传甲.txt"}
    old = store.snapshot(first["library_id"], first["generation"])
    assert len(old["documents"]) == 1
    assert store.document(saved["documents"][1]["id"], saved["library_id"], saved["generation"])[1][0]["text"] == "上传原文"


def test_failed_and_cancelled_scans_do_not_switch_or_replace(store, tmp_path):
    (store.default_root / "good.txt").write_text("原来的可用证据", encoding="utf-8")
    before = build(store)
    other = tmp_path / "bad"
    other.mkdir()
    (other / "bad.docx").write_bytes(b"invalid zip")
    with pytest.raises(RuntimeError, match="解析失败"):
        build(store, other)
    assert store.status()["generation"] == before["generation"]
    cancellation = threading.Event()
    def cancel_during_read(fields):
        if "message" in fields:
            cancellation.set()
    with pytest.raises(Cancelled):
        store.build(store.default_root, cancel_during_read, cancellation)
    assert store.status()["generation"] == before["generation"]
    with pytest.raises(ValueError):
        build(store, tmp_path / "missing")
    assert store.active_id == before["library_id"]


def test_duplicate_bytes_keep_both_filenames_and_private_folders_are_pruned(store):
    for name in ("原始 甲.txt", "原始 乙.txt"):
        (store.default_root / name).write_text("相同内容两份来源", encoding="utf-8")
    hidden = store.default_root / "node_modules"
    hidden.mkdir()
    (hidden / "不可解析.docx").write_bytes(b"not a zip")
    build(store)
    docs = store.snapshot()["documents"]
    assert {d["name"] for d in docs} == {"原始 甲.txt", "原始 乙.txt"}
    assert len({d["id"] for d in docs}) == 2
    assert len({d["blob"] for d in docs}) == 1


def test_remove_keeps_original_and_old_citation_version(store):
    path = store.default_root / "保留.txt"
    path.write_text("原文件不删除", encoding="utf-8")
    before = build(store)
    doc = store.snapshot()["documents"][0]
    store.delete(doc["id"], before["generation"])
    assert path.exists()
    assert not store.snapshot()["documents"]
    assert store.document(doc["id"], before["library_id"], before["generation"])[1]
    with pytest.raises(ValueError):
        store.delete(doc["id"], before["generation"])


def test_json_compatibility_record_failure_does_not_turn_commit_into_failed_job(store):
    (store.default_root / "a.txt").write_text("available", encoding="utf-8")
    with patch("jev_core.storage.atomic_json", side_effect=OSError("read-only record")):
        result = build(store)
    assert store.status()["generation"] == result["generation"]


class FakeModel:
    def __init__(self, reject=False):
        self.reject = reject
        self.evaluated = 0
    def rank(self, question, candidates):
        self.evaluated = len(candidates)
        return [{**item, "jev": {"probability": .5, "none_probability": .1,
                 "supports_over_none": not self.reject, "choice": "none" if self.reject else "c0",
                 "batch": i//4}} for i, item in enumerate(candidates)]
    def status(self):
        return {"local_only": True}


def test_all_requested_candidates_and_only_selected_citations(store):
    (store.default_root / "文章.txt").write_text("\n\n".join(f"检索测试 证据条目{i}" for i in range(16)), encoding="utf-8")
    build(store)
    model = FakeModel()
    result = RagService(store, model).query("检索测试", 12, 2)
    assert model.evaluated == 12
    assert len(result["evidence"]) == len(result["citations"]) == 2
    for evidence, citation in zip(result["evidence"], result["citations"]):
        assert evidence["text"] == citation["text"]
        assert evidence["block_id"] == citation["block_id"]
        assert evidence["generation"] == result["generation"]


def test_one_candidate_can_be_rejected_and_empty_library_is_distinct(store):
    service = RagService(store, FakeModel(reject=True))
    assert service.query("你好")["status"] == "empty"
    (store.default_root / "a.txt").write_text("股票行情内容", encoding="utf-8")
    build(store)
    result = service.query("股票行情")
    assert result["status"] == "no_answer"
    assert result["citations"] == result["evidence"] == []
    assert result["diagnostics"]["local_jev_evaluated"] == 1


def test_cjk_tokens_do_not_join_unrelated_punctuation():
    assert "乙丙" not in tokens("甲乙。丙丁")
    assert tokens("甲乙。丙丁") == ["甲乙", "丙丁"]
    assert "是什" not in query_tokens("资料室客服邮箱是什么？")


def test_shared_topic_without_the_requested_fact_is_not_returned(store):
    (store.default_root / "行情.txt").write_text("股票行情数据更新延迟不超过 500ms。", encoding="utf-8")
    build(store)
    result = RagService(store, FakeModel()).query("股票行情供应商的 API 密钥是多少？")
    assert result["status"] == "no_answer"


def test_picker_cancel_and_static_fallback_contract(store, monkeypatch):
    import app as server
    from fastapi.testclient import TestClient
    monkeypatch.setattr(server, "store", store)
    monkeypatch.setattr(server, "choose_directory", lambda: "")
    monkeypatch.setattr(server, "dist", store.default_root / "no-build")
    with TestClient(server.app, base_url="http://127.0.0.1:8765") as client:
        assert client.post("/api/source-root/pick").json() == {"cancelled": True, "path": None}
        assert store.snapshot()["generation"] is None
        assert client.get("/static/app.js").status_code == 200
        assert client.get("/").status_code == 200
        assert isinstance(client.get("/api/documents").json()["documents"], list)
        assert client.post("/api/query", json={"question":"  "}).status_code == 422
        assert client.get("/api/config", headers={"Origin":"https://unrelated.example"}).status_code == 403
        assert client.get("/api/config", headers={"Host":"unrelated.example"}).status_code == 400


def test_api_upload_preserves_unicode_original_and_source_switch_isolates_it(store, tmp_path, monkeypatch):
    import time
    import app as server
    from fastapi.testclient import TestClient
    from jev_core.storage import IndexJobs
    monkeypatch.setattr(server, "store", store)
    monkeypatch.setattr(server, "jobs", IndexJobs(store))
    monkeypatch.setattr(server, "rag_service", RagService(store, FakeModel()))
    def completed(client):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            job = client.get("/api/index/job").json()["job"]
            if job["state"] != "running":
                assert job["state"] == "completed", job
                return
            time.sleep(.02)
        pytest.fail("index job did not finish")
    with TestClient(server.app, base_url="http://127.0.0.1:8765") as client:
        original = "原始上传金额 876.54 元。\n".encode()
        response = client.post("/api/documents/upload", files=[("files", ("资料 原文.txt", original, "text/plain"))])
        assert response.status_code == 202
        completed(client)
        data = client.get("/api/documents").json()
        doc, status = data["documents"][0], data["status"]
        assert doc["name"] == "资料 原文.txt"
        link = f"/api/documents/{doc['id']}/raw?library_id={status['library_id']}&generation={status['generation']}"
        assert client.get(link).content == original
        query = client.post("/api/query", json={"question":"原始上传金额"}).json()
        assert query["status"] == "found"
        assert "876.54" in query["evidence"][0]["text"]
        other = tmp_path / "second"
        other.mkdir()
        assert client.post("/api/source-root", json={"path":str(other)}).status_code == 202
        completed(client)
        assert not client.get("/api/documents").json()["documents"]
        assert client.get(link).content == original
