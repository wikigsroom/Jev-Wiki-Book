"""Exercise the actual native binary: auth, publication, query boundary, CPU JEV."""
import argparse
import http.cookiejar
import json
import os
from pathlib import Path
import signal
import socket
import sqlite3
import subprocess
import tempfile
import time
import urllib.error
import urllib.request


def free_port():
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        return sock.getsockname()[1]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--binary", type=Path, required=True)
    parser.add_argument("--models", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--ui", action="store_true", help="Also test admin and query forms in CI Chromium")
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workspace = Path(tempfile.mkdtemp(prefix="native-smoke-", dir=args.output.parent))
    docs = workspace / "documents"
    docs.mkdir()
    (docs / "资料室.md").write_text("# 资料室\n\n资料室开放时间为每周二至周六 09:30–17:30。\n", encoding="utf-8")
    (docs / "private.md").write_text("绝密代号是 SECRET-NOT-PUBLISHED-429。", encoding="utf-8")
    # Create fixtures with native text and an embedded image. The frozen runtime
    # must parse DOCX/PDF and actually execute OCR, not just import its modules.
    from PIL import Image, ImageDraw, ImageFont
    from docx import Document
    import pymupdf
    picture = Image.new("RGB", (1000, 150), "white")
    font_path = "C:/Windows/Fonts/arial.ttf" if os.name == "nt" else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    font = ImageFont.truetype(font_path, 48)
    ImageDraw.Draw(picture).text((20, 35), "JEVOFFLINETEST2026", fill="black", font=font)
    image_file = workspace / "ocr.png"
    picture.save(image_file)
    word = Document()
    word.add_paragraph("DOCX native parsing verification.")
    word.add_picture(str(image_file))
    word.save(docs / "format-check.docx")
    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "PDF native parsing verification.")
        pdf.save(docs / "format-check.pdf")
    admin_port, web_port = free_port(), free_port()
    while web_port == admin_port:
        web_port = free_port()
    origin = f"http://127.0.0.1:{admin_port}"
    query_origin = f"http://127.0.0.1:{web_port}"
    opener = urllib.request.build_opener(urllib.request.ProxyHandler({}), urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    csrf = ""

    def request(url, data=None):
        payload = json.dumps(data).encode() if data is not None else None
        headers = {"Content-Type": "application/json", "X-CSRF-Token": csrf}
        try:
            with opener.open(urllib.request.Request(url, data=payload, headers=headers), timeout=180) as response:
                return response.status, json.loads(response.read())
        except urllib.error.HTTPError as error:
            return error.code, json.loads(error.read())

    log = (workspace / "server.log").open("w", encoding="utf-8")
    command = [str(args.binary.resolve()), "--data-dir", str(workspace / "data"), "--models-dir", str(args.models.resolve()),
               "--admin-port", str(admin_port), "--web-port", str(web_port), "--web-host", "127.0.0.1"]

    def start():
        process = subprocess.Popen(command, stdout=log, stderr=log, creationflags=subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0)
        for _ in range(240):
            try:
                if request(query_origin + "/api/health")[0] == 200:
                    return process
            except OSError:
                pass
            if process.poll() is not None:
                raise RuntimeError("Binary exited during startup: " + (workspace / "server.log").read_text(encoding="utf-8"))
            time.sleep(.5)
        else:
            process.kill()
            process.wait(timeout=10)
            raise RuntimeError("Binary never became ready")

    def stop(process):
        process.send_signal(signal.CTRL_BREAK_EVENT if os.name == "nt" else signal.SIGTERM)
        assert process.wait(timeout=30) == 0
        for port in [admin_port, web_port]:
            with socket.socket() as probe:
                assert probe.connect_ex(("127.0.0.1", port)) != 0

    process = start()
    try:
        duplicate = subprocess.run(command, capture_output=True, text=True, timeout=30)
        assert duplicate.returncode != 0 and "already in use" in duplicate.stderr
        missing = subprocess.run([str(args.binary.resolve()), "--models-dir", str(workspace / "missing-models"), "--check"], capture_output=True, text=True, timeout=60)
        assert missing.returncode == 2 and not json.loads(missing.stdout)["ready"]
        assert request(origin + "/api/admin/libraries")[0] == 401
        code, session = request(origin + "/api/admin/login", {"username": "admin", "password": "Jev-Change-Me-2026!"})
        assert code == 200
        csrf = session["csrf_token"]
        assert request(origin + "/api/admin/status")[0] == 403
        code, session = request(origin + "/api/admin/password", {"current_password": "Jev-Change-Me-2026!", "new_password": "Native-Smoke-Only-2026!"})
        assert code == 200
        csrf = session["csrf_token"]
        assert request(origin + "/api/admin/scan", {"path": str(docs)})[0] == 202
        for _ in range(120):
            state = request(origin + "/api/admin/status")[1]
            if state["job"]["state"] != "running":
                assert state["job"]["state"] == "completed", state
                break
            time.sleep(.5)
        libraries = request(origin + "/api/admin/libraries")[1]["libraries"]
        library = next(item for item in libraries if item["documents"])
        # Admin owns local files; these snapshots are never served to query users.
        with sqlite3.connect(workspace / "data/storage/knowledge.sqlite3") as database:
            snapshots = database.execute("SELECT payload FROM snapshots").fetchall()
        blocks = [block for (payload,) in snapshots for block in json.loads(payload)["blocks"]]
        content = "\n".join(block["text"] for block in blocks)
        assert "DOCX native parsing verification." in content, content
        assert "PDF native parsing verification." in content, content
        assert "JEVOFFLINETEST2026" in content.replace(" ", ""), content
        document = next(item for item in library["documents"] if item["name"] == "资料室.md")
        grant = {"library_id": library["id"], "generation": library["generation"], "document_ids": [document["id"]]}
        assert request(origin + "/api/admin/publication", {"grants": [grant]})[0] == 200
        started = time.monotonic()
        code, answer = request(query_origin + "/api/query", {"question": "资料室开放时间是什么？"})
        query_seconds = round(time.monotonic() - started, 3)
        assert code == 200 and answer["evidence"] and "09:30" in answer["evidence"][0]["text"], answer
        assert str(docs) not in json.dumps(answer)
        assert "SECRET-NOT-PUBLISHED-429" not in json.dumps(answer)
        for route in ["/api/admin/status", "/api/documents", "/api/source-root", "/api/config"]:
            assert request(query_origin + route)[0] == 404
        model = request(origin + "/api/admin/status")[1]["model"]
        assert model["loaded"] and model["device"] == "cpu" and model["precision"] == "fp32", model
        revision = request(origin + "/api/admin/status")[1]["publication"]["revision"]
        stop(process)
        process = start()  # Same ports and same persisted data, without a delay.
        assert request(origin + "/api/admin/status")[0] == 401
        assert request(origin + "/api/admin/login", {"username": "admin", "password": "Jev-Change-Me-2026!"})[0] == 401
        code, session = request(origin + "/api/admin/login", {"username": "admin", "password": "Native-Smoke-Only-2026!"})
        assert code == 200
        csrf = session["csrf_token"]
        assert request(origin + "/api/admin/status")[1]["publication"]["revision"] == revision
        code, restored = request(query_origin + "/api/query", {"question": "资料室开放时间是什么？"})
        assert code == 200 and "09:30" in restored["evidence"][0]["text"]
        assert request(origin + "/api/admin/publication", {"grants": []})[0] == 200
        assert request(query_origin + "/api/query", {"question": "开放时间"})[0] == 503
        if args.ui:
            from ui_smoke import verify_ui
            verify_ui(origin, query_origin, docs, args.output.parent / "ui-verification")
        report = {"passed": True, "binary": args.binary.name, "model": model, "query_seconds": query_seconds,
                  "auth_checked": True, "password_changed": True, "publication_checked": True, "private_routes_blocked": True, "revocation_checked": True,
                  "docx_checked": True, "pdf_checked": True, "ocr_checked": True, "ui_checked": args.ui,
                  "immediate_restart_same_ports": True, "persisted_password_and_grants": True, "old_session_revoked_on_restart": True,
                  "data_directory_lock": True, "missing_models_exit_nonzero": True}
        stop(process)
        args.output.write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(json.dumps(report), flush=True)
    finally:
        if process.poll() is None:
            if os.name == "nt":
                process.send_signal(signal.CTRL_BREAK_EVENT)
            else:
                process.terminate()
            try:
                process.wait(timeout=15)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=10)
        log.close()


if __name__ == "__main__":
    main()
